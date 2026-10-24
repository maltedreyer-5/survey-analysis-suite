# SPDX-License-Identifier: MIT
# SPDX-FileCopyrightText: 2026 Malte Dreyer
"""numeric_matrix handler: statistics per field, constant-sum check."""
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from handlers import get_handler  # noqa: E402


def test_statistics_per_field_and_constant_sum():
    df = pd.DataFrame({"a": ["50", "60", "30", None], "b": ["50", "30", "70", None], "c": ["0", "10", "0", None]})
    h = get_handler("numeric_matrix", "budget", {"id": "budget", "columns": ["a", "b", "c"], "constant_sum": 100,
                                       "item_labels": {"de": ["Personal", "Sachmittel", "Sonstiges"],
                                                       "en": ["Staff", "Material", "Other"]}})
    r = h.analyze(df, lang="en")
    assert r.n == 3 and list(r.data) == ["Staff", "Material", "Other"]
    assert r.data["Staff"]["mean"] == 46.7 and r.data["Other"]["max"] == 10
    cs = r.metadata["constant_sum"]
    assert cs["n_matching"] == 3 and cs["share_matching"] == 100.0
    assert round(cs["mean_share_of_total"]["Material"], 1) == 50.0
    assert r.description == '"Material" has the highest mean (50.0; n = 3).'


def test_sums_that_do_not_match_are_reported():
    df = pd.DataFrame({"a": [50, 20], "b": [50, 20]})
    r = get_handler("numeric_matrix", "x", {"id": "x", "columns": ["a", "b"], "constant_sum": 100}).analyze(df, lang="de")
    assert r.metadata["constant_sum"]["n_matching"] == 1 and any("50%" in w for w in r.warnings)


def test_registered_and_validated():
    from services.config_validation import validate_survey_config
    cfg = {"settings": {"survey_metadata": {"source_language": "de"}},
           "question_groups": {"x": {"handler": "numeric_matrix", "columns": ["a", "b"]}}}
    errors, warnings = validate_survey_config(cfg)
    assert errors == [] and not any("noch keine eigene Darstellung" in w for w in warnings)   # display exists (part 2)


def test_word_table_and_test_variables(tmp_path):
    from docx import Document
    from utils.analysis_variables import build_variables
    from utils.export_word import WordExporter
    df = pd.DataFrame({"a": [50, 60, 30], "b": [50, 40, 70]})
    qcfg = {"id": "budget", "handler": "numeric_matrix", "columns": ["a", "b"], "constant_sum": 100,
            "item_labels": {"de": ["Personal", "Sachmittel"]}, "label": {"de": "Budget"}}
    r = get_handler("numeric_matrix", "budget", qcfg).analyze(df, lang="de")
    doc = Document()
    WordExporter(tmp_path)._add_numeric_matrix_table(doc, r.metadata["item_stats"], r.metadata["constant_sum"], "de")
    cells = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert cells[:2] == ["Feld", "n"] and "Personal" in cells and "Anteil an der Summe" in cells
    assert any("kommen auf die Summe 100" in p.text for p in doc.paragraphs)
    cfg = {"settings": {"survey_metadata": {"source_language": "de"}}, "question_groups": {"budget": qcfg}}
    variables = build_variables(cfg, df, ["budget"])["budget"]
    assert [v.level for v in variables] == ["numeric", "numeric"] and variables[0].values.tolist() == [50, 60, 30]


def test_dashboard_format(monkeypatch):
    import sys as _sys
    import types
    from export.dashboard import _enrich_numeric_matrix_question
    from export.result_format import _format_result_for_dashboard
    cfg = {"settings": {"survey_metadata": {"source_language": "de", "languages": ["de", "en"]}}}
    monkeypatch.setitem(_sys.modules, "core", types.SimpleNamespace(state=types.SimpleNamespace(config=cfg)))
    qcfg = {"id": "budget", "columns": ["a", "b"], "constant_sum": 100,
            "item_labels": {"de": ["Personal", "Sachmittel"], "en": ["Staff"]}}
    r = get_handler("numeric_matrix", "budget", qcfg).analyze(pd.DataFrame({"a": [50, 60], "b": [50, 40]}), lang="de")
    formatted = _format_result_for_dashboard({"n": r.n, "data": r.data, "metadata": r.metadata,
                                                                         "percentages": None}, "numeric_matrix")
    assert set(formatted["fields"]) == {"f0", "f1"} and formatted["constant_sum"]["mean_share"]["f1"] == 45.0
    q = _enrich_numeric_matrix_question({"id": "budget"}, qcfg, ["de", "en"])
    assert q["items"][0]["label"] == {"de": "Personal", "en": "Staff"} and q["items"][1]["label"]["en"] == "Sachmittel"
    assert q["constant_sum"] == 100.0


def test_summary_data_section():
    from utils.question_summary import DATA_SECTION_BUILDERS
    df = pd.DataFrame({"a": [50, 60], "b": [50, 40]})
    qcfg = {"id": "b", "columns": ["a", "b"], "constant_sum": 100, "item_labels": {"de": ["Personal", "Sachmittel"]}}
    r = get_handler("numeric_matrix", "b", qcfg).analyze(df, lang="de")
    text = DATA_SECTION_BUILDERS["numeric_matrix"](r, qcfg, "de")
    assert "mehrere Zahlenfelder (n=2)" in text and "Personal: n=2" in text and "Konstantsumme 100: 100 %" in text


def test_blank_cells_are_no_answers_in_every_handler():
    df = pd.DataFrame({"x": ["A", "A", "B", None, "", "  "]})
    r = get_handler("single_choice", "sc", {"id": "sc", "columns": ["x"]}).analyze(df, lang="de")
    assert r.n == 3 and dict(r.data) == {"A": 2, "B": 1}
    n = get_handler("numeric", "nu", {"id": "nu", "columns": ["v"]}).analyze(
        pd.DataFrame({"v": ["1", "2", "abc", None, ""]}), lang="de")
    assert n.n == 2 and n.metadata["n_invalid"] == 1            # only "abc" is invalid; blanks are no answer


def test_cooperation_matrix_default_scale_and_unknown_values():
    df = pd.DataFrame({"K [A]": ["Sehr wichtig", "Wichtig", "Keine Kooperation"], "K [B]": ["Keine Kooperation"] * 3})
    r = get_handler("cooperation_matrix", "k", {"id": "k", "columns": ["K [A]", "K [B]"]}).analyze(df, lang="de")
    scores = r.metadata["importance_scores"]
    assert scores["A"] > scores["B"]                         # "no cooperation" is the lowest, not the highest
    odd = get_handler("cooperation_matrix", "k", {"id": "k", "columns": ["K [A]"]}).analyze(
        pd.DataFrame({"K [A]": ["Zentral", "Wichtig"]}), lang="de")
    assert any("outside the scale" in w and "Zentral" in w for w in odd.warnings)
