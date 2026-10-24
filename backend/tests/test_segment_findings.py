# SPDX-License-Identifier: MIT
# SPDX-FileCopyrightText: 2026 Malte Dreyer
"""
Differences between groups (utils/segment_findings.py): same tests as the
summaries, selection, facts, fixed sentences, free-text topics through the
global clusters, texts in the source language with a separate possible
reading, stored texts.
"""
import json
import pathlib
import sys
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from utils.analysis_texts import AnalysisTextStore  # noqa: E402
from utils.persistence import SurveyDatabase  # noqa: E402
from utils.segment_findings import SegmentDifferenceAnalyzer  # noqa: E402

RNG = np.random.default_rng(11)


def survey(lang="de", n=80, two_findings=False):
    seg = np.array(["A"] * n + ["B"] * n)
    df = pd.DataFrame({"seg": seg})
    # option 0: large difference; option 1: significant with many cases but small gap; option 2: none
    p = {"A": (0.2, 0.50, 0.4), "B": (0.8, 0.62, 0.4)}
    if two_findings:   # option 2 differs as well, the other way round
        p = {"A": (0.2, 0.50, 0.8), "B": (0.8, 0.62, 0.15)}
    for k in range(3):
        df[f"o{k}"] = [("x" if RNG.random() < p[s][k] else None) for s in seg]
    df["o3"] = "x"   # chosen by everybody, so the base is all respondents and shares are unconditional
    df["noise"] = RNG.choice(["a", "b"], len(df))
    df["open"] = ["text"] * len(df)
    labels = {"de": ("Gruppe", "Dienste", "Zufall", "Wünsche"), "en": ("Group", "Services", "Random", "Wishes")}[lang]
    cfg = {"settings": {"survey_metadata": {"source_language": lang, "languages": [lang, "fr"]},
                        "llm_context": "Campus",
                        "segmentation": {"dimensions": [
                            {"id": "grp", "label": {lang: labels[0]}, "source_question": "seg",
                             "bucketing": {"type": "identity"}}]}},
           "question_groups": {
               "seg": {"handler": "single_choice", "columns": ["seg"], "label": {lang: labels[0]}},
               "mc": {"handler": "multi_choice", "columns": ["o0", "o1", "o2", "o3"], "label": {lang: labels[1]},
                      "item_labels": {lang: ["Mensa", "Bibliothek", "Sport", "WLAN"]}},
               "noise": {"handler": "single_choice", "columns": ["noise"], "label": {lang: labels[2]}},
               "open": {"handler": "freetext", "columns": ["open"], "label": {lang: labels[3]}}}}
    return cfg, df


def clusters_for(df, lang="de"):
    """Analysis result of the free-text question: topic 'Parken' mostly from group B."""
    b = df.index[df["seg"] == "B"]
    a = df.index[df["seg"] == "A"]
    parking = set(b[: int(len(b) * 0.7)]) | set(a[: int(len(a) * 0.1)])
    other = set(df.index) - parking
    clusters = [SimpleNamespace(name={lang: "Parken"}, response_ids=parking),
                SimpleNamespace(name={lang: "Sonstiges"}, response_ids=other)]
    return {"open": SimpleNamespace(clustering_result=SimpleNamespace(clusters=clusters))}


class FakeLLM:
    def __init__(self):
        self.calls = []

    def complete(self, prompt, system_prompt=None, **kw):
        self.calls.append((system_prompt or "", prompt))
        if prompt.rstrip().endswith("]"):                       # translation batch
            texts = json.loads(prompt[prompt.rindex("\n[") + 1:])
            return json.dumps([f"[fr] {t}" for t in texts], ensure_ascii=False)
        if "Gesamtbild" in (system_prompt or "") or "overall picture" in (system_prompt or ""):
            return "Gesamtbild."
        return json.dumps({"description": "Beschreibung.", "reading": "Lesart."})


def test_large_difference_listed_small_gap_counted_unrelated_absent():
    cfg, df = survey(n=400)
    report = SegmentDifferenceAnalyzer(cfg, df).run(languages=["de"])
    dim = report.questions["mc"].dimensions[0]
    listed = [f.result.var_b.id for f in dim.findings]
    assert listed == ["mc::0"]                       # 60 points between the groups
    assert dim.small_differences == 1                 # option 1: significant, gap about 12 points
    assert dim.state == "differs"
    shares = {g["value"]: round(g["share"]) for g in dim.findings[0].groups}
    assert shares["B"] > 70 and shares["A"] < 30


def test_same_tests_as_summary():
    from utils.question_summary import compute_segment_significance
    cfg, df = survey()
    report = SegmentDifferenceAnalyzer(cfg, df).run(languages=["de"])
    summary = compute_segment_significance("mc", "multi_choice", cfg["question_groups"]["mc"], df, cfg)
    ours = report.questions["mc"].dimensions[0]
    assert summary[0].state == ours.state
    assert ours.findings[0].result.q_value == pytest.approx(summary[0].findings[0].q_value)


def test_fixed_sentences_for_dimensions_without_listed_findings():
    cfg, df = survey()
    report = SegmentDifferenceAnalyzer(cfg, df).run(languages=["de", "en"])
    data = report.to_json(cfg, ["de", "en"])
    noise = data["questions"]["noise"]["dimensions"]["grp"]
    assert noise["state"] == "no_difference" and not noise["findings"]
    assert "keine statistisch signifikanten Unterschiede" in noise["statement"]["de"]
    assert "Gruppe" in noise["statement"]["de"]
    small = SegmentDifferenceAnalyzer(cfg, df.iloc[list(range(8)) + list(range(80, 88))]).run(languages=["de"])
    assert small.questions["mc"].dimensions[0].state == "not_testable"
    assert "seg" not in data["questions"] or not data["questions"]["seg"]["dimensions"]  # own segmentation


def test_free_text_topics_tested_through_global_clusters():
    cfg, df = survey()
    report = SegmentDifferenceAnalyzer(cfg, df, analysis_results=clusters_for(df)).run(languages=["de"])
    findings = report.questions["open"].dimensions[0].findings
    assert findings and findings[0].result.var_b.measure == "cluster_mentioned"
    facts = findings[0].facts("de", "de")
    assert any("Parken" in f and "je Gruppe" in f for f in facts)
    # without clusters the question is not listed at all (nothing could be tested)
    assert "open" not in SegmentDifferenceAnalyzer(cfg, df).run(languages=["de"]).questions


def test_summary_segment_part_includes_free_text_topics():
    from utils.analysis_variables import cluster_memberships_from_results
    from utils.question_summary import compute_segment_significance
    cfg, df = survey()
    memberships = cluster_memberships_from_results(clusters_for(df), cfg)
    sig = compute_segment_significance("open", "freetext", cfg["question_groups"]["open"], df, cfg,
                                       cluster_memberships=memberships)
    assert sig and sig[0].state == "differs"


@pytest.mark.parametrize("lang,marker,question", [("de", "Deutsch", "Frage: Dienste"),
                                                  ("en", "English", "Question: Services")])
def test_texts_from_catalog_of_source_language_then_translated(lang, marker, question):
    cfg, df = survey(lang)
    llm = FakeLLM()
    report = SegmentDifferenceAnalyzer(cfg, df, llm_client=llm).run(languages=[lang, "fr"])
    finding = report.questions["mc"].dimensions[0].findings[0]
    system, user = llm.calls[0]
    assert marker in system and "Campus" in system and '"reading"' in system
    assert question in user and ("Fakten:" in user or "Facts:" in user)
    assert finding.description == {lang: "Beschreibung.", "fr": "[fr] Beschreibung."}
    assert finding.reading == {lang: "Lesart.", "fr": "[fr] Lesart."}


def test_without_llm_only_facts():
    cfg, df = survey()
    data = SegmentDifferenceAnalyzer(cfg, df).run(languages=["de", "fr"]).to_json(cfg, ["de", "fr"])
    f = data["questions"]["mc"]["dimensions"]["grp"]["findings"][0]
    assert f["description"] == {} and f["reading"] == {}
    assert set(f["facts"]) == {"de", "fr"} and f["q_value"] is not None
    assert data["$schema"] == "segment_insights_v2" and data["method"]["correction"] == "benjamini_hochberg"
    json.dumps(data)


def test_stored_texts_reused_and_manual_kept(tmp_path):
    cfg, df = survey()
    db = SurveyDatabase(tmp_path / "s.db")
    with db._connection() as conn:
        conn.execute("INSERT INTO projects (id, name) VALUES (1, 'p')")
        conn.execute("INSERT INTO imports (id, project_id, filename) VALUES (1, 1, 'f.csv')")
    store = AnalysisTextStore(db, 1, "segment_difference")
    SegmentDifferenceAnalyzer(cfg, df, llm_client=FakeLLM(), text_store=store, db=db).run(languages=["de"])
    second = FakeLLM()
    report = SegmentDifferenceAnalyzer(cfg, df, llm_client=second, text_store=store, db=db).run(languages=["de"])
    assert not [c for c in second.calls if '"reading"' in c[0]]
    key = report.questions["mc"].dimensions[0].findings[0].key
    store.put(key, "x", {"de": "Von Hand."}, {}, manual=True)
    third = SegmentDifferenceAnalyzer(cfg, df, llm_client=FakeLLM(), text_store=store, db=db).run(languages=["de"])
    assert third.questions["mc"].dimensions[0].findings[0].description == {"de": "Von Hand."}


def test_configuration_error_is_reported_not_skipped():
    cfg, df = survey()
    cfg["settings"]["segmentation"]["dimensions"][0]["source_question"] = "does_not_exist"
    report = SegmentDifferenceAnalyzer(cfg, df).run(languages=["de"])
    assert report.errors


def _doc_text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def test_word_section_shows_facts_description_and_reading(tmp_path):
    from docx import Document
    from utils.export_word import WordExporter
    cfg, df = survey()
    report = SegmentDifferenceAnalyzer(cfg, df, analysis_results=clusters_for(df),
                                       llm_client=FakeLLM()).run(languages=["de", "fr"])
    doc = Document()
    WordExporter(tmp_path).add_group_differences_section(doc, report, cfg, "de")
    body = _doc_text(doc)
    assert "Teil III: Unterschiede zwischen Gruppen" in body and "15 Prozentpunkten" in body
    assert "Nach Gruppe" in body and "Anteil „Dienste: Mensa“ je Gruppe" in body
    assert "Beschreibung." in body and "Mögliche Lesart: Lesart." in body
    cfg2, df2 = survey(two_findings=True)
    report2 = SegmentDifferenceAnalyzer(cfg2, df2, llm_client=FakeLLM()).run(languages=["de"])
    doc2 = Document()
    WordExporter(tmp_path).add_group_differences_section(doc2, report2, cfg2, "de")
    assert "Zusammenfassung: Gesamtbild." in _doc_text(doc2)
    assert "Parken" in body                                     # free-text topic tested
    assert "Zufall" not in body                                 # questions without differences are not listed


def test_word_segment_topics_are_descriptive(tmp_path):
    from docx import Document
    from utils.export_word import WordExporter
    data = {"dimension_label": {"de": "Region"}, "questions": {"q": {
        "question_label": {"de": "Wünsche"},
        "clusters_by_value": {"nord": {"segment_value_label": {"de": "Nord"}, "total_responses": 20,
                                      "clusters": [{"name": {"de": "Parken"}, "count": 9, "percentage": 45.0}]}}}}}
    doc = Document()
    WordExporter(tmp_path).add_segment_comparison_section(doc, data, "de")
    body = _doc_text(doc)
    assert "Teil IV: Themen je Gruppe nach Region" in body
    assert "nicht getestet" in body and "Nord (n = 20)" in body and "Parken: 45 %" in body


def test_segment_clusters_are_counted_by_the_handler(monkeypatch):
    """segment_clusters.json takes counts from the free-text handler, never from LLM estimates."""
    import export.freetext_segments as fs
    from core import state
    from export.insights import export_segment_cluster_comparison
    cfg, df = survey()
    cfg["settings"]["segmentation"]["dimensions"][0].update(enable_specific_clusters=True,
                                                            specific_clusters_min_n=10)
    calls = []

    def fake_specific(qid, qconfig, dim, value):
        calls.append((qid, dim, value))
        return {"n": 80, "clusters": [{"name": {"de": "Parken"}, "unique_respondents": 20, "count": 25},
                                      {"name": {"de": "Mensa"}, "unique_respondents": 40, "count": 41}]}
    monkeypatch.setattr(fs, "_specific_clusters", fake_specific)
    monkeypatch.setattr(state, "df", df)
    monkeypatch.setattr(state, "config", cfg)
    data = export_segment_cluster_comparison(None, ["de"])
    groups = data["dimensions"]["grp"]["questions"]["open"]["clusters_by_value"]
    assert sorted(calls) == [("open", "grp", "A"), ("open", "grp", "B")]
    first = next(iter(groups.values()))
    assert [c["name"]["de"] for c in first["clusters"]] == ["Mensa", "Parken"]      # largest first
    assert first["clusters"][0]["percentage"] == 50.0 and "comparison_insight" not in data["dimensions"]["grp"]["questions"]["open"]


def test_specific_cluster_cache_only_within_a_run(monkeypatch):
    import export.freetext_segments as fs
    from core import state
    n = []
    monkeypatch.setattr(fs, "_cluster_segment", lambda *a: n.append(a) or {"n": 1})
    monkeypatch.setattr(state, "import_id", 1)
    fs._specific_clusters("q", {}, "d", "v"); fs._specific_clusters("q", {}, "d", "v")
    assert len(n) == 2                       # no cache outside a pipeline run
    fs.begin_specific_cluster_cache()
    try:
        fs._specific_clusters("q", {}, "d", "v"); fs._specific_clusters("q", {}, "d", "v")
    finally:
        fs.end_specific_cluster_cache()
    assert len(n) == 3


def test_dimension_summary_from_two_findings_stored_and_protected(tmp_path):
    cfg, df = survey(two_findings=True)
    db = SurveyDatabase(tmp_path / "s.db")
    with db._connection() as conn:
        conn.execute("INSERT INTO projects (id, name) VALUES (1, 'p')")
        conn.execute("INSERT INTO imports (id, project_id, filename) VALUES (1, 1, 'f.csv')")
    store = AnalysisTextStore(db, 1, "segment_difference")
    llm = FakeLLM()
    report = SegmentDifferenceAnalyzer(cfg, df, llm_client=llm, text_store=store, db=db).run(languages=["de", "fr"])
    dim = report.questions["mc"].dimensions[0]
    assert len(dim.findings) == 2
    assert dim.summary == {"de": "Gesamtbild.", "fr": "[fr] Gesamtbild."}
    system, user = next(c for c in llm.calls if "Gesamtbild" in c[0])
    assert "Dienste" in system and "Gruppe" in system and "Campus" in system
    assert user.count("- Beschreibung.") == 2 and "%" not in user        # descriptions only, no facts
    noise = report.questions["noise"].dimensions[0]
    assert noise.summary == {}                                            # nothing listed, no summary
    second = FakeLLM()
    again = SegmentDifferenceAnalyzer(cfg, df, llm_client=second, text_store=store, db=db).run(languages=["de", "fr"])
    assert not [c for c in second.calls if "Gesamtbild" in c[0]]
    assert again.questions["mc"].dimensions[0].summary["fr"] == "[fr] Gesamtbild."
    store.put("summary:mc|grp", "x", {"de": "Von Hand."}, {}, manual=True)
    third = SegmentDifferenceAnalyzer(cfg, df, llm_client=FakeLLM(), text_store=store, db=db).run(languages=["de"])
    assert third.questions["mc"].dimensions[0].summary["de"] == "Von Hand."
    data = third.to_json(cfg, ["de"])
    assert data["questions"]["mc"]["dimensions"]["grp"]["summary"]["de"] == "Von Hand."
    assert data["parameters"]["summary_min_findings"] == 2


def test_single_finding_gets_no_summary():
    cfg, df = survey()
    llm = FakeLLM()
    report = SegmentDifferenceAnalyzer(cfg, df, llm_client=llm).run(languages=["de"])
    assert len(report.questions["mc"].dimensions[0].findings) == 1
    assert report.questions["mc"].dimensions[0].summary == {}
    assert not [c for c in llm.calls if "Gesamtbild" in c[0]]
