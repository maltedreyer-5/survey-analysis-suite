# SPDX-License-Identifier: MIT
# SPDX-FileCopyrightText: 2026 Malte Dreyer
"""Publication documents: paths are independent of the working directory; the cover has
neutral placeholders; a survey without documents.yaml gets a clear message."""
import os
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from utils import document_builder as dbm  # noqa: E402


def make(tmp_path, extra=""):
    survey = tmp_path / "survey"; (survey / "docs").mkdir(parents=True)
    (survey / "docs" / "cover.md").write_text("# {title}\n\n{organization} · {year} · {doi_url}\n", encoding="utf-8")
    (survey / "documents.yaml").write_text(
        'source_language: "en"\nsource_dir: "docs"\n'
        'metadata: {survey_title: {en: "Title"}, organization: "Org", year: 2026}\n'
        'doi: {en: "https://doi.org/10.1/x"}\nshared_sections: {cover: "cover.md"}\n'
        'documents:\n  rep:\n    filename: rep\n    existing_docx: "run_*/report_{lang}_*.docx"\n' + extra, encoding="utf-8")
    return survey


def test_source_dir_is_relative_to_the_documents_yaml(tmp_path, monkeypatch):
    survey = make(tmp_path)
    monkeypatch.chdir(tmp_path)                     # any working directory
    b = dbm.DocumentBuilder(config_path=str(survey / "documents.yaml"), db_path=str(tmp_path / "t.db"))
    assert b.source_dir == survey / "docs"
    cover = b._resolve_cover(b.config["documents"]["rep"], "en")
    assert cover.startswith("# Title") and "Org · 2026 · https://doi.org/10.1/x" in cover


def test_existing_docx_is_relative_to_the_output_dir_and_takes_lang(tmp_path, monkeypatch):
    survey = make(tmp_path)
    out = tmp_path / "outputs"; (out / "run_1").mkdir(parents=True)
    old = out / "run_1" / "report_it_1.docx"; old.write_bytes(b"x")
    new = out / "run_1" / "report_it_2.docx"; new.write_bytes(b"y")
    os.utime(old, (1, 1))
    monkeypatch.setattr(dbm, "_output_dir", lambda: out)
    b = dbm.DocumentBuilder(config_path=str(survey / "documents.yaml"), db_path=str(tmp_path / "t.db"))
    pattern = b._existing_docx_pattern(b.config["documents"]["rep"], "it")
    assert pattern == "run_*/report_it_*.docx"
    assert b._resolve_existing_docx(pattern) == new          # the newest match
    assert b._existing_docx_pattern({"existing_docx": {"de": "a.docx"}}, "de") == "a.docx"


def test_survey_without_documents_yaml_gets_a_clear_message(tmp_path):
    with pytest.raises(FileNotFoundError, match="publication documents are optional"):
        dbm.DocumentBuilder(config_path=str(tmp_path / "documents.yaml"), db_path=str(tmp_path / "t.db"))


def test_existing_docx_only_from_runs_of_the_active_survey(tmp_path, monkeypatch):
    import json, os
    survey = make(tmp_path)
    out = tmp_path / "outputs"
    for run, sid, t in (("run_1", "this_survey", 1), ("run_2", "other_survey", 2), ("run_3", None, 0)):
        (out / run).mkdir(parents=True)
        f = out / run / "report_en_1.docx"; f.write_bytes(b"x"); os.utime(f, (t, t))
        if sid:
            (out / run / "dashboard_data.json").write_text(json.dumps({"metadata": {"survey_id": sid}}))
    monkeypatch.setattr(dbm, "_output_dir", lambda: out)
    monkeypatch.setattr(dbm, "_active_survey_id", lambda: "this_survey")
    b = dbm.DocumentBuilder(config_path=str(survey / "documents.yaml"), db_path=str(tmp_path / "t.db"))
    # the newest report belongs to another survey: the newest of this survey is taken
    assert b._resolve_existing_docx("run_*/report_en_*.docx") == out / "run_1" / "report_en_1.docx"


def test_paragraphs_heading_followed_by_text():
    paras = dbm.DocumentBuilder._paragraphs("# A\nText a\n\n## B\nText b\n\n```\n# code\n```")
    assert paras == ["# A", "Text a", "## B", "Text b", "```\n# code\n```"]


@pytest.mark.parametrize("part,toc", [("Part I: Question Analyses", "Table of Contents"),
                                      ("Parte I: Analisi delle domande", "Indice"),
                                      ("Teil I: Fragenauswertungen", "Inhaltsverzeichnis")])
def test_cover_and_toc_of_a_report_are_removed_in_every_language(tmp_path, part, toc):
    from docx import Document
    from docx.enum.text import WD_BREAK
    survey = make(tmp_path)
    d = Document()
    d.add_paragraph("Survey title", style="Title")
    d.add_paragraph("October 2026").add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph(toc, style="TOC Heading")
    d.add_paragraph("").add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph(part, style="Heading 1")
    d.add_paragraph("Question 1", style="Heading 2")
    d.add_paragraph("Figure 1")
    b = dbm.DocumentBuilder(config_path=str(survey / "documents.yaml"), db_path=str(tmp_path / "t.db"))
    stripped = b._strip_cover_and_toc(d)
    texts = [p.text for p in stripped.paragraphs]
    assert texts == [part, "Question 1", "Figure 1"]
