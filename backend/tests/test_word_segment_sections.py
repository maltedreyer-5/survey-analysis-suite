# SPDX-License-Identifier: MIT
# SPDX-FileCopyrightText: 2026 Malte Dreyer
"""Word report sections use configured dimensions and catalog texts."""
import sys
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import core  # noqa: E402
from export import segment_context  # noqa: E402
from utils.export_word import WordExporter  # noqa: E402

CONFIG = {
    "settings": {"survey_metadata": {"languages": ["de", "it"]},
                 "segmentation": {"dimensions": [
                     {"id": "region", "label": {"de": "Region", "it": "Regione"}, "source_question": "q_region",
                      "value_labels": {"Süd": {"it": "Sud"}}}]}},
    "question_groups": {"q_region": {"handler": "single_choice", "columns": ["R"]}},
}


@pytest.fixture
def exporter(monkeypatch, tmp_path):
    monkeypatch.setattr(core.state, "config", CONFIG)
    monkeypatch.setattr(core.state, "df", pd.DataFrame({"R": ["Nord", "Süd", "Süd"]}))
    segment_context.reset_context()
    yield WordExporter(output_dir=tmp_path) if "output_dir" in WordExporter.__init__.__code__.co_varnames else WordExporter(tmp_path)
    segment_context.reset_context()


def texts(doc):
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return [p.text for p in doc.paragraphs] + cells


def test_distribution_table_for_any_dimension(exporter):
    doc = Document()
    exporter._add_segment_table(doc, [("Süd", 2), ("Nord", 1)], 3, "it", "region")
    t = texts(doc)
    assert "Distribuzione per Regione" in t
    assert "Sud" in t and "Regione" in t and "Quota" in t


def test_segment_theme_details_use_value_labels(exporter):
    doc = Document()
    result = SimpleNamespace(metadata={"cluster_details": [{"name": "Costi", "count": 2}]}, data={}, percentages={})
    exporter._add_segment_theme_details(doc, "region", {"Süd": {"result": result, "n": 2}}, "it")
    t = texts(doc)
    assert "Sud (n=2)" in t and any("Costi" in x for x in t)
