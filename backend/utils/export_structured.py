# SPDX-License-Identifier: MIT
# SPDX-FileCopyrightText: 2026 Malte Dreyer
"""
Structured Word export of survey reports, following the question groups of the
survey with charts and tables.
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Any
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import matplotlib
from utils.languages import text as _text
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from .charts import ChartGenerator
from .cluster_examples import select_cluster_examples
from .translation import t_report, get_translation_manager

logger = logging.getLogger(__name__)


def _src() -> str:
    """Source language of the loaded survey (labels in the data are in this language)."""
    from utils.prompts import source_language
    return source_language()



class StructuredReportExporter:
    """
    Creates structured Word reports by question group.
    
    Structure:
    1. title page
    2. table of contents
    3. main part: results by question group
       - per question: overall result + segment charts
    4. appendix A: detailed tables with significance
    5. appendix B: free-text clusters with examples
    """
    
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.chart_gen = ChartGenerator()
        self.figure_counter = 0
        self.table_counter = 0
        self._translation_cache = {}  # cache of the category translations
        self._llm = None
        self._tm = get_translation_manager()
    
    def _translate_segment_name(self, seg_name: str, segment_type: str, lang: str) -> str:
        """
        Translates a segment name with TranslationManager.get_segment_label(),
        i.e. from the labels of the segmentation configuration.
        """
        if lang == _src() or not self._tm:
            return seg_name
        return self._tm.get_segment_label(segment_type, seg_name, lang)
    
    def _add_help_text(self, doc, help_text: str):
        """
        Adds an explanation (help text) after the question title.
        
        Args:
            doc: Word document
            help_text: the explanation
        """
        if not help_text:
            return
        
        from docx.shared import Pt, RGBColor
        para = doc.add_paragraph()
        run = para.add_run(help_text)
        run.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(100, 100, 100)  # grey
        para.space_after = Pt(8)
    
    def _get_category_translation(self, category: str, qconfig: Dict, lang: str) -> str:
        """
        Translates a category into the target language: first from the options
        of the configuration, then from the catalogs (common answers).
        """
        from utils.languages import get_source_language
        from core import state
        src = get_source_language(state.config)
        if lang == src:
            return category
        
        # check the cache
        cache_key = f"{category}:{lang}"
        if cache_key in self._translation_cache:
            return self._translation_cache[cache_key]
        
        # from the configuration (options in the source language -> options in lang)
        options = qconfig.get('options', {})
        if isinstance(options, dict):
            de_options = options.get(src, [])  # options in the source language
            target_options = options.get(lang, [])
            
            if de_options and target_options and len(de_options) == len(target_options):
                mapping = dict(zip(de_options, target_options))
                if category in mapping:
                    translated = mapping[category]
                    self._translation_cache[cache_key] = translated
                    return translated
        
        # common answers from the catalogs (e.g. "Sonstiges" -> "Autres")
        if self._tm:
            tm_result = self._tm.get_category_translation(category, lang)
            if tm_result != category:
                self._translation_cache[cache_key] = tm_result
                return tm_result
        
        # LLM fallback through the single translation path (memory, context)
        try:
            from utils.translation_service import get_translation_service, context_for_question
            qid = qconfig.get('id') or next((k for k, v in (state.config or {}).get('question_groups', {}).items()
                                              if v is qconfig), '')
            translated = get_translation_service().translate(
                category, lang, context_for_question(state.config, qid, 'option'))
            if translated:
                self._translation_cache[cache_key] = translated
                return translated
        except Exception as e:
            logger.warning(f"Translation of the category failed: {e}")
        
        # fallback: return the original
        return category
    
    def _translate_data_keys(self, data: Dict, qconfig: Dict, lang: str) -> Dict:
        """Translates the keys of a data dictionary."""
        # a DataFrame (matrix) is left as it is: the handler translates its labels when drawing
        if lang == _src() or _is_empty(data) or hasattr(data, "empty"):
            return data
        
        return {
            self._get_category_translation(k, qconfig, lang): v
            for k, v in data.items()
        }
    
    def create_report(
        self,
        all_results: Dict,
        config: Dict,
        title: str = "Umfrage-Auswertung",
        subtitle: str = "",
        lang: str = "de",
        include_small_multiples: bool = True,
        include_grouped_segments: bool = True,
        n_cluster_examples: int = 5
    ) -> Document:
        """
        Creates the structured report.
        
        Args:
            all_results: results of analyze_all_questions_with_segments()
            config: configuration with question_groups and report_groups
            title: report title
            subtitle: subtitle
            lang: language
            include_small_multiples: include small multiples
            include_grouped_segments: include grouped segment charts
            n_cluster_examples: number of cluster examples
        
        Returns:
            Word document
        """
        from handlers import get_handler
        from handlers.base import AnalysisResult
        
        doc = Document()
        self.figure_counter = 0
        self.table_counter = 0
        
        # report groups from the configuration
        settings = config.get('settings', {})
        report_groups = settings.get('report_groups', {})
        question_groups = config.get('question_groups', {})
        
        # === 1. TITLE PAGE ===
        self._add_title_page(doc, title, subtitle, lang)
        
        # === 2. TABLE OF CONTENTS ===
        self._add_toc_placeholder(doc, lang)
        
        # === 3. MAIN PART: RESULTS BY GROUP ===
        doc.add_heading(t_report("main.title", lang), level=1)
        
        # groups sorted by order
        sorted_groups = sorted(
            report_groups.items(),
            key=lambda x: x[1].get('order', 99)
        )
        
        # collect tables and clusters for the appendices
        all_tables = []
        all_clusters = []
        
        for group_id, group_config in sorted_groups:
            group_label = group_config.get('label', {}).get(lang, group_id)
            group_description = group_config.get('description', {}).get(lang, '')
            group_questions = group_config.get('questions', [])
            
            # group heading
            doc.add_heading(group_label, level=2)
            
            if group_description:
                desc_para = doc.add_paragraph()
                desc_para.add_run(group_description).italic = True
                desc_para.space_after = Pt(12)
            
            # questions of this group
            for qid in group_questions:
                if qid not in all_results:
                    continue
                
                q_data = all_results[qid]
                if 'error' in q_data:
                    continue
                
                qconfig = question_groups.get(qid, {})
                question_text = q_data.get('question_text', qid)
                handler_type = q_data.get('handler_type', 'single_choice')
                total_result = q_data.get('total')
                by_segment = q_data.get('by_segment', {})
                
                # question heading
                doc.add_heading(question_text, level=3)
                
                # help text (explanation) after the title
                from export.question_texts import help_text as _help_text
                help_text = _help_text(qconfig, lang)
                if help_text:
                    self._add_help_text(doc, help_text)
                
                # overall result
                if total_result:
                    self._add_question_result(
                        doc, qid, qconfig, total_result, handler_type, lang
                    )
                
                # segment charts
                for seg_key in by_segment.keys():  # configured dimensions, in configured order
                    seg_data = by_segment.get(seg_key, {})
                    segments = seg_data.get('segments', {})
                    
                    if not segments:
                        continue
                    
                    seg_label = seg_data.get('label', seg_key)
                    
                    # a question is not compared with the groups formed from its own answers
                    if _dimension_source(seg_key) == qid:
                        continue
                    
                    if handler_type in ('numeric', 'numeric_matrix'):
                        if handler_type == 'numeric':
                            self._add_numeric_segment_chart(doc, question_text, segments, seg_label, lang, qconfig)
                    else:
                        # grouped segment chart
                        if include_grouped_segments:
                            self._add_grouped_segment_chart(
                                doc, question_text, segments, seg_label, lang, qconfig,
                                segment_type=seg_key
                            )
                        
                        # Small Multiples
                        if include_small_multiples:
                            self._add_small_multiples_chart(
                                doc, question_text, segments, seg_label, lang, qconfig,
                                segment_type=seg_key
                            )
                    
                    # collect the table for the appendix
                    all_tables.append({
                        'qid': qid,
                        'question_text': question_text,
                        'seg_key': seg_key,
                        'seg_label': seg_label,
                        'segments': segments,
                        'qconfig': qconfig
                    })
                
                # collect the free-text clusters for the appendix
                if handler_type == 'freetext' and total_result:
                    metadata = self._get_metadata(total_result)
                    cluster_details = metadata.get('cluster_details', [])
                    if cluster_details:
                        all_clusters.append({
                            'qid': qid,
                            'question_text': question_text,
                            'clusters': cluster_details
                        })
                
                doc.add_paragraph()  # spacing
        
        doc.add_page_break()
        
        # === 4. APPENDIX A: DETAILED TABLES ===
        self._add_appendix_tables(doc, all_tables, lang)
        
        # === 5. APPENDIX B: FREE-TEXT CLUSTERS ===
        if all_clusters:
            self._add_appendix_clusters(doc, all_clusters, lang, n_cluster_examples)
        
        return doc
    
    def _add_title_page(self, doc: Document, title: str, subtitle: str, lang: str):
        """Adds the title page."""
        # title
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(28)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_para.space_before = Pt(100)
        title_para.space_after = Pt(24)
        
        # subtitle
        if subtitle:
            subtitle_para = doc.add_paragraph()
            subtitle_run = subtitle_para.add_run(subtitle)
            subtitle_run.font.size = Pt(16)
            subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # date
        date_str = datetime.now().strftime("%d.%m.%Y")
        date_label = t_report("title_page.date_label", lang)
        date_para = doc.add_paragraph()
        date_run = date_para.add_run(f"{date_label} {date_str}")
        date_run.font.size = Pt(11)
        date_run.italic = True
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_page_break()
    
    def _add_toc_placeholder(self, doc: Document, lang: str):
        """Adds the table of contents placeholder."""
        doc.add_heading(t_report("toc.title", lang), level=1)
        
        para = doc.add_paragraph()
        para.add_run(_text("word.toc_placeholder", lang)).italic = True
        
        doc.add_page_break()
    
    def _add_question_result(
        self,
        doc: Document,
        qid: str,
        qconfig: Dict,
        result: Any,
        handler_type: str,
        lang: str
    ):
        """Adds the overall result of a question."""
        from handlers import get_handler
        from handlers.base import AnalysisResult
        
        # extract the data
        if hasattr(result, 'n'):
            n = result.n
            data = result.data
            percentages = getattr(result, 'percentages', {})
            metadata = getattr(result, 'metadata', {})
        else:
            n = result.get('n', 0)
            data = result.get('data', {})
            percentages = result.get('percentages', {})
            metadata = result.get('metadata', {})
        
        if _is_empty(data):
            return
        
        # translate the categories if needed (reports in other languages)
        if lang != _src() and qconfig:
            data = self._translate_data_keys(data, qconfig, lang)
            if percentages:
                percentages = self._translate_data_keys(percentages, qconfig, lang)
        
        # result object for the handler (with translated data)
        result_obj = AnalysisResult(
            question_id=qid,
            n=n,
            data=data,
            percentages=percentages,
            metadata=metadata
        )
        
        try:
            handler = get_handler(handler_type, qid, qconfig)
            if handler:
                default_chart = qconfig.get('default_chart', 'bar_horizontal')
                fig = handler.create_chart(result_obj, default_chart, lang=lang)
                
                if fig:
                    self._add_figure(doc, fig, n, lang)
                    plt.close(fig)
        except Exception as e:
            logger.warning(f"Chart creation failed for {qid}: {e}")
    
    def _add_matrix_segment_charts(self, doc, question_text, chart_data, seg_label, lang, qconfig):
        """
        Matrix question by segment, as set in matrix_segment_chart (question, else
        settings, default both): small multiples per row and/or the mean per row and segment.
        """
        from core import state
        mode = matrix_segment_chart_mode(qconfig, (state.config or {}).get('settings', {}))
        if mode == 'none':
            return
        scale = matrix_scale(qconfig, chart_data, lang)
        if mode in ('both', 'small_multiples'):
            fig = self.chart_gen.matrix_rows_by_segment(chart_data, scale, title=_chart_title(question_text, seg_label),
                                                        segment_label=seg_label, lang=lang,
                                                        colors=self._scale_colours(qconfig, lang))
            if fig:
                self._add_figure(doc, fig, None, lang, caption=_text("word.by_segment", lang, segment=seg_label))
                plt.close(fig)
        if mode in ('both', 'mean'):
            means = matrix_means(chart_data, scale, qconfig)
            if not means:
                logger.debug(f"Matrix means skipped (no scale direction): {question_text[:40]}")
                return
            fig = self.chart_gen.matrix_means_by_segment(means, title=_chart_title(question_text, seg_label),
                                                         segment_label=seg_label, lang=lang)
            if fig:
                self._add_figure(doc, fig, None, lang, caption=_text("word.by_segment", lang, segment=seg_label))
                plt.close(fig)

    def _scale_colours(self, qconfig: Dict, lang: str) -> Dict:
        """Colours of the answers of an ordered scale (utils/chart_colors.py), in the labels of ``lang``."""
        from utils.chart_colors import scale_colours
        src = _src()
        for key in ('options', 'scale'):
            value = qconfig.get(key)
            if isinstance(value, dict) and value.get(src):
                source = list(value[src])
                shown = list(value.get(lang) or source)
                colours = scale_colours(source, qconfig)
                if len(shown) == len(source):
                    return {shown[i]: colours[o] for i, o in enumerate(source) if o in colours}
                return colours
        return {}

    def _add_numeric_segment_chart(self, doc, question_text, segments, seg_label, lang, qconfig):
        """Numeric question by group: median and quartiles per group (the tests compare distributions)."""
        stats = {}
        for seg_name, seg_data in segments.items():
            result = seg_data.get('result')
            data = result.data if hasattr(result, 'data') else (result or {}).get('data', {})
            if isinstance(data, dict) and data.get('median') is not None:
                stats[seg_name] = {'mean': data['median'], 'low': data.get('q25'), 'high': data.get('q75'),
                                   'n': data.get('n', seg_data.get('n', 0))}
        if not stats:
            return
        try:
            fig = self.chart_gen.segment_means_bar(
                segment_stats=stats, title=_chart_title(question_text, seg_label),
                xlabel=str((qconfig.get('numeric') or {}).get('unit') or ''),
                decimals=int((qconfig.get('numeric') or {}).get('decimals', 1)),
                note=_text("chart.medians", lang, dimension=seg_label))
            if fig:
                self._add_figure(doc, fig, None, lang, caption=_text("word.by_segment", lang, segment=seg_label))
                plt.close(fig)
        except Exception as e:
            logger.warning(f"Numeric chart by group failed: {e}")

    def _add_grouped_segment_chart(
        self,
        doc: Document,
        question_text: str,
        segments: Dict,
        seg_label: str,
        lang: str,
        qconfig: Dict = None,
        segment_type: str = ""
    ):
        """Adds the grouped segment chart."""
        # prepare the data for the grouped bar chart
        chart_data = {}
        
        for seg_name, seg_data in segments.items():
            result = seg_data.get('result')
            if not result:
                continue
            
            # extract the percentages
            if hasattr(result, 'percentages'):
                percentages = result.percentages or {}
            elif isinstance(result, dict):
                percentages = result.get('percentages', {})
            else:
                continue
            
            if percentages:
                # translate the categories if needed
                from core import state
                from utils.languages import get_source_language
                if qconfig and lang != get_source_language(state.config):
                    percentages = self._translate_data_keys(percentages, qconfig, lang)
                # translate the segment name (e.g. Student -> Étudiant·e)
                display_name = self._translate_segment_name(seg_name, segment_type, lang)
                chart_data[display_name] = percentages
        
        if not chart_data:
            return
        # a matrix has a share per row and category (a dict per category): its own charts
        if is_matrix_data(chart_data):
            self._add_matrix_segment_charts(doc, question_text, chart_data, seg_label, lang, qconfig or {})
            return
        
        try:
            from services.question_analysis import _category_order
            fig = self.chart_gen.bar_grouped(
                category_order=_category_order(qconfig or {}, lang),
                data=chart_data,
                title=_chart_title(question_text, seg_label),
                segment_label=seg_label,
                show_percentages=True,
                lang=lang,
                figsize=(10, 6)
            )
            
            if fig:
                self._add_figure(doc, fig, None, lang, caption=_text("word.by_segment", lang, segment=seg_label))
                plt.close(fig)
        except Exception as e:
            logger.warning(f"Grouped chart failed: {e}")
    
    def _add_small_multiples_chart(
        self,
        doc: Document,
        question_text: str,
        segments: Dict,
        seg_label: str,
        lang: str,
        qconfig: Dict = None,
        segment_type: str = ""
    ):
        """Adds the small multiples chart."""
        try:
            # translate the categories if needed
            translated_segments = segments
            if qconfig and lang != _src():
                translated_segments = {}
                for seg_name, seg_data in segments.items():
                    new_seg_data = dict(seg_data)  # copy
                    result = seg_data.get('result')
                    if result:
                        if hasattr(result, 'percentages') and result.percentages:
                            # copy of the result with translated keys
                            from handlers.base import AnalysisResult
                            new_result = AnalysisResult(
                                question_id=getattr(result, 'question_id', ''),
                                n=result.n,
                                data=self._translate_data_keys(result.data if not _is_empty(result.data) else {}, qconfig, lang),
                                percentages=self._translate_data_keys(result.percentages or {}, qconfig, lang),
                                description=result.description,
                                warnings=result.warnings,
                                clusters=result.clusters,
                                metadata=result.metadata
                            )
                            new_seg_data['result'] = new_result
                        elif isinstance(result, dict):
                            new_result = dict(result)
                            if 'percentages' in new_result:
                                new_result['percentages'] = self._translate_data_keys(
                                    new_result['percentages'], qconfig, lang
                                )
                            if 'data' in new_result:
                                new_result['data'] = self._translate_data_keys(
                                    new_result['data'], qconfig, lang
                                )
                            new_seg_data['result'] = new_result
                    translated_segments[seg_name] = new_seg_data
            
            # translate the segment names (e.g. Student -> Étudiant·e)
            if lang != _src():
                translated_segments = {
                    self._translate_segment_name(k, segment_type, lang): v
                    for k, v in translated_segments.items()
                }
            
            title = _chart_title(question_text, _text("word.question_comparison", lang, question="", segment=seg_label).strip(" -–:"))
            from services.question_analysis import _category_order
            fig = self.chart_gen.small_multiples(
                segment_results=translated_segments,
                title=title,
                segment_label=seg_label,
                show_percentages=True,
                lang=lang,
                max_cols=3,
                category_order=_category_order(qconfig or {}, lang),
                colors=self._scale_colours(qconfig or {}, lang),
            )
            
            if fig:
                self._add_figure(doc, fig, None, lang, caption=_text("word.small_multiples_by", lang, segment=seg_label))
                plt.close(fig)
        except Exception as e:
            logger.warning(f"Small multiples failed: {e}")
    
    def _add_figure(
        self,
        doc: Document,
        fig: plt.Figure,
        n: Optional[int],
        lang: str,
        caption: str = None,
        width: float = 6.0
    ):
        """Adds a figure with a caption."""
        self.figure_counter += 1
        
        # convert the figure to bytes
        buf = BytesIO()
        fig.savefig(buf, format='png', bbox_inches='tight', dpi=150)
        buf.seek(0)
        
        # close the figure to avoid a memory leak
        plt.close(fig)
        
        # insert the image
        doc.add_picture(buf, width=Inches(width))
        
        # caption
        caption_text = f"{_text('word.figure', lang).strip()} {self.figure_counter}"
        if caption:
            caption_text += f": {caption}"
        if n:
            caption_text += f" (n={n})"
        
        caption_para = doc.add_paragraph()
        caption_para.add_run(caption_text).italic = True
        caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_para.space_after = Pt(12)
    
    def _add_appendix_tables(self, doc: Document, all_tables: List[Dict], lang: str):
        """Adds appendix A with the detailed tables."""
        doc.add_heading(t_report("appendix_a.title", lang), level=1)
        
        for table_info in all_tables:
            qid = table_info['qid']
            question_text = table_info['question_text']
            seg_label = table_info['seg_label']
            segments = table_info['segments']
            qconfig = table_info.get('qconfig', {})
            segment_type = table_info.get('seg_key', '')
            
            # heading
            doc.add_heading(f"{question_text} - {seg_label}", level=3)
            
            # Significance: per answer option / item, adjusted (utils/segment_tests.py)
            line = self._segment_test_line(qid, segment_type, lang)
            if line:
                sig_para = doc.add_paragraph()
                sig_para.add_run(f"{t_report('statistics.significance', lang)}: ").bold = True
                sig_para.add_run(line)
                sig_para.space_after = Pt(6)
            
            # create the table
            self._add_comparison_table(doc, segments, lang, qconfig, segment_type=segment_type)
            
            doc.add_paragraph()  # spacing
        
        doc.add_page_break()
    
    def _segment_test_line(self, qid: str, dimension_id: str, lang: str) -> str:
        """Result of the segment tests of one question and dimension, as one line."""
        from core import state
        from utils.correlation import finding_facts
        from utils.languages import get_source_language, localized, text
        from utils.segment_tests import run_segment_tests
        if state.df is None or not state.config:
            return ""
        key = (qid, id(state.df))
        cache = getattr(self, "_segment_tests_cache", {})
        if key not in cache:
            try:
                from utils.analysis_variables import cluster_memberships_from_results
                memberships = cluster_memberships_from_results(state.analysis_results or {}, state.config)
                cache[key] = run_segment_tests(state.config, state.df, qid, memberships)
            except (ValueError, KeyError) as e:
                logger.warning(f"Segment tests for {qid} not possible: {e}")
                cache[key] = None
            self._segment_tests_cache = cache
        tests = cache[key]
        dim = next((d for d in (tests.dimensions if tests else []) if d.dimension_id == dimension_id), None)
        if dim is None:
            return ""
        src = get_source_language(state.config)
        label = str(localized(dim.label, lang, src) or dimension_id)
        if dim.state == "not_testable":
            return text("summary.segments_not_testable", lang, dimensions=label)
        if dim.state == "no_difference":
            from utils.question_summary import no_significant_differences_text
            return no_significant_differences_text([label], lang)
        return " ".join(finding_facts(r, lang, src)[0] + " — " + "; ".join(finding_facts(r, lang, src)[1:])
                        for r in dim.significant()[:2])

    def _add_comparison_table(self, doc: Document, segments: Dict, lang: str,
                              qconfig: Dict = None, segment_type: str = ""):
        """Adds a comparison table."""
        self.table_counter += 1
        
        if not segments:
            return
        
        # collect all categories
        all_categories = set()
        for seg_data in segments.values():
            result = seg_data.get('result')
            if result:
                data = result.data if hasattr(result, 'data') else result.get('data', {})
                if isinstance(data, dict):
                    all_categories.update(data.keys())
        
        categories = sorted(all_categories, key=str)
        if not categories:
            return
        
        # create the table
        segment_names = list(segments.keys())
        n_cols = len(segment_names) + 1
        n_rows = len(categories) + 2  # header + n row + categories
        
        table = doc.add_table(rows=n_rows, cols=n_cols)
        table.style = 'Table Grid'
        
        # header (translate the segment names)
        header_row = table.rows[0]
        header_row.cells[0].text = t_report("appendix_a.category", lang)
        for i, seg_name in enumerate(segment_names):
            n = segments[seg_name].get('n', 0)
            warning = segments[seg_name].get('warning', '')
            display_name = self._translate_segment_name(seg_name, segment_type, lang)
            cell_text = display_name
            if warning:
                cell_text += " ⚠️"
            header_row.cells[i + 1].text = cell_text
        
        # n row
        n_row = table.rows[1]
        n_row.cells[0].text = "n"
        for i, seg_name in enumerate(segment_names):
            n = segments[seg_name].get('n', 0)
            n_row.cells[i + 1].text = str(n)
        
        # data rows
        for row_idx, category in enumerate(categories):
            row = table.rows[row_idx + 2]
            # translate the category if needed
            display_category = category
            if qconfig and lang != _src():
                display_category = self._get_category_translation(category, qconfig, lang)
            row.cells[0].text = str(display_category)[:50]  # shorten
            
            for col_idx, seg_name in enumerate(segment_names):
                result = segments[seg_name].get('result')
                if result:
                    data = result.data if hasattr(result, 'data') else result.get('data', {})
                    percentages = result.percentages if hasattr(result, 'percentages') else result.get('percentages', {})
                    
                    count = data.get(category, 0) if isinstance(data, dict) else 0
                    pct = percentages.get(category, 0) if isinstance(percentages, dict) else 0
                    
                    row.cells[col_idx + 1].text = f"{count} ({pct:.1f}%)"
                else:
                    row.cells[col_idx + 1].text = "-"
    
    def _add_appendix_clusters(
        self,
        doc: Document,
        all_clusters: List[Dict],
        lang: str,
        n_examples: int = 5
    ):
        """Adds appendix B with the free-text clusters."""
        from core import state
        from utils.languages import get_source_language
        from utils.translation_service import get_translation_service, context_for_question
        
        doc.add_heading(t_report("appendix_b.title", lang), level=1)
        
        # Translations only for languages other than the source language
        src = get_source_language(state.config)
        translator = get_translation_service() if lang != src else None
        
        for cluster_info in all_clusters:
            question_text = cluster_info['question_text']
            clusters = cluster_info['clusters']
            
            # Question text, cluster names, descriptions and examples of one
            # question are translated together, with the question as context
            qid = cluster_info.get('qid', '')
            if translator:
                question_text = translator.translate(
                    question_text, lang, context_for_question(state.config, qid, 'question'))
            
            doc.add_heading(question_text, level=2)
            
            # choose the best examples with the LLM
            try:
                clusters = select_cluster_examples(clusters, n_examples, lang)
            except Exception as e:
                logger.warning(f"Selection of the cluster examples failed: {e}")
            
            for cluster in clusters:
                name = cluster.get('name') or _text('word.unnamed', lang)
                description = cluster.get('description', '')
                count = cluster.get('count', 0)
                selected = cluster.get('selected_examples', cluster.get('examples', []))
                
                if translator:
                    name, description = translator.translate_many(
                        [name, description], lang,
                        context_for_question(state.config, qid, 'cluster_name'))
                
                # cluster heading
                cluster_para = doc.add_paragraph()
                cluster_para.add_run(f"{name}").bold = True
                cluster_para.add_run(f" ({count} {t_report('appendix_b.mentions', lang)})")
                
                if description:
                    desc_para = doc.add_paragraph()
                    desc_para.add_run(description).italic = True
                
                # translate and show the examples
                if selected:
                    doc.add_paragraph().add_run(
                        t_report("appendix_b.selected_examples", lang)
                    ).bold = True
                    
                    shown = list(selected[:n_examples])
                    if translator:
                        shown = translator.translate_many(
                            shown, lang, context_for_question(state.config, qid, 'answer'))
                    for example in shown:

                        example_para = doc.add_paragraph(style='List Bullet')
                        example_para.add_run(f'"{example}"')
                
                doc.add_paragraph()  # spacing
    
    def _to_analysis_result(self, result) -> Any:
        """Converts a dict into an AnalysisResult-like object."""
        from handlers.base import AnalysisResult
        
        if hasattr(result, 'n'):
            return result
        
        return AnalysisResult(
            question_id=result.get('question_id', ''),
            n=result.get('n', 0),
            data=result.get('data', {}),
            percentages=result.get('percentages', {}),
            description=result.get('description', ''),
            warnings=result.get('warnings', []),
            clusters=result.get('clusters', []),
            metadata=result.get('metadata', {})
        )
    
    def _get_metadata(self, result) -> Dict:
        """Extracts the metadata from a result."""
        if hasattr(result, 'metadata'):
            return result.metadata or {}
        elif isinstance(result, dict):
            return result.get('metadata', {})
        return {}
    
    def export_report(
        self,
        all_results: Dict,
        config: Dict,
        filename: str = "survey_report",
        **kwargs
    ) -> Path:
        """
        Creates and saves the report.
        
        Args:
            all_results: results of analyze_all_questions_with_segments()
            config: configuration
            filename: file name (without extension)
            **kwargs: further arguments for create_report
        
        Returns:
            path of the created file
        """
        lang = kwargs.get('lang') or _src()
        doc = self.create_report(all_results, config, **kwargs)
        
        filepath = self.output_dir / f"{filename}_{lang}.docx"
        doc.save(filepath)
        logger.info(f"Report saved: {filepath}")
        
        return filepath


def export_structured_report(
    all_results: Dict,
    config: Dict,
    output_dir: Path,
    filename: str = "survey_report",
    **kwargs
) -> Path:
    """
    Main function of the structured report export.
    
    Args:
        all_results: results
        config: configuration
        output_dir: output directory
        filename: file name
        **kwargs: further options
    
    Returns:
        path of the file
    """
    exporter = StructuredReportExporter(output_dir)
    return exporter.export_report(all_results, config, filename, **kwargs)

def _chart_title(question: str, grouping: str, width: int = 70) -> str:
    """Title of a chart by group: the question (shortened at a word boundary) and, on a line
    of its own, the grouping in full."""
    import textwrap
    q = textwrap.shorten(str(question), width=width, placeholder=" …")
    return f"{q}\n{grouping}" if grouping else q


def _dimension_source(dimension_id: str):
    """The question a segmentation dimension is formed from (None for a data column)."""
    try:
        from core import state
        from utils.segment_resolver import SegmentResolver
        dim = SegmentResolver(state.config or {}).get_dimension(dimension_id)
        return getattr(dim, 'source_question', None)
    except Exception:
        return None


def _is_empty(data) -> bool:
    """True for None, an empty dict/list or an empty DataFrame (a DataFrame has no truth value)."""
    if data is None:
        return True
    if hasattr(data, "empty"):
        return bool(data.empty)
    return not data


MATRIX_SEGMENT_CHART_MODES = ("both", "small_multiples", "mean", "none")


def matrix_segment_chart_mode(qconfig: dict, settings: dict) -> str:
    """matrix_segment_chart of the question, else of settings, else 'both' (unknown values: 'both')."""
    mode = (qconfig or {}).get("matrix_segment_chart") or (settings or {}).get("matrix_segment_chart") or "both"
    return mode if mode in MATRIX_SEGMENT_CHART_MODES else "both"


def is_matrix_data(chart_data: dict) -> bool:
    """{segment: {category: value}} where the values are dicts (a share per matrix row)."""
    return any(isinstance(v, dict) for pct in chart_data.values() for v in (pct or {}).values())


def matrix_scale(qconfig: dict, chart_data: dict, lang: str = None) -> list:
    """
    Scale options of a matrix question in the language of the data: the configured
    scale ({lang: [...]} or a list) whose options occur in chart_data, else the
    categories in the data in their order.
    """
    in_data = list(dict.fromkeys(o for pct in chart_data.values() for o in pct))
    scale = qconfig.get('scale')
    candidates = []
    if isinstance(scale, dict):
        candidates = [scale.get(lang)] + list(scale.values()) if lang else list(scale.values())
    elif scale:
        candidates = [scale]
    for cand in candidates:
        opts = [str(x) for x in (cand or [])]
        if opts and set(opts) & set(map(str, in_data)):
            return opts
    return in_data


def matrix_means(chart_data: dict, scale: list, qconfig: dict) -> dict:
    """{segment: {row: mean on the scale}} from {segment: {option: {row: percent}}} (empty without a scale direction)."""
    from utils.scale_metrics import compute
    means = {}
    for seg, by_option in chart_data.items():
        rows = list(dict.fromkeys(r for opt in by_option.values() for r in (opt or {})))
        for row in rows:
            weights = {str(o): int(round(float((by_option.get(o) or {}).get(row) or 0) * 10)) for o in by_option}
            metrics = compute(weights, scale, qconfig)
            if metrics and metrics.get('mean') is not None:
                means.setdefault(seg, {})[row] = metrics['mean']
    return means
